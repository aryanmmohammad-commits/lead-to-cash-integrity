import csv
import re
import shutil
import xml.etree.ElementTree as ET
from dataclasses import replace
from decimal import Decimal
from html.parser import HTMLParser
from pathlib import Path
from unittest.mock import Mock

import duckdb
import pytest
import requests
from openpyxl import load_workbook
from typer.testing import CliRunner

from l2c import cli
from l2c.fixes import propose_fixes
from l2c.output_data import Table, chf, read_review_data
from l2c.report import owner_groups, quarantine_summary, render_html, waterfall, write_report
from l2c.slack import build_digest, format_digest, send_digest


class ReportParser(HTMLParser):
    def __init__(self, text: str) -> None:
        super().__init__()
        self.tags: list[str] = []
        self.references: list[str] = []
        self.feed(text)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        self.tags.append(tag)
        for key, value in attrs:
            if key in ("src", "href", "action") and value and not value.startswith("#"):
                self.references.append(value)


@pytest.mark.parametrize("format_name", ["default", "ch", "de", "hubspot_ui"])
def test_report_and_csv_contracts_for_every_export_format(
    built: dict[str, Path], format_name: str, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("L2C_DB_PATH", str(built[format_name]))
    report = write_report(output_dir=tmp_path)
    proposal = propose_fixes(output_dir=tmp_path)
    text = report.html_path.read_text(encoding="utf-8")
    assert re.search(r'data-record-id="PLANT-R04".*?CHF 18,000\.00', text, re.DOTALL)
    assert "CHF 826,224.00" in text and "CHF 736,008.00" in text and "CHF -90,216.00" in text
    assert "Synthetic data. Independent demo." in text
    assert "not valued" in text and "no_fx_rate" in text
    assert "change line" not in text and "score history" not in text
    parser = ReportParser(text)
    assert not parser.references
    assert not {"script", "link", "iframe", "img"}.intersection(parser.tags)
    assert "@import" not in text and "url(" not in text
    svg = ET.fromstring(re.search(r"<svg\b.*?</svg>", text, re.DOTALL).group())
    bars = svg.findall(".//{http://www.w3.org/2000/svg}rect")
    data = read_review_data(built[format_name])
    assert len(bars) == sum(row["amount_chf"] != 0 for row in data.tables["arr_bridge"].rows) + 2
    assert bars[0].get("class") == bars[-1].get("class") == "total"
    svg_text = " ".join(svg.itertext())
    assert "Start: CHF 0" in svg_text and "Total gap: CHF -90,216" in svg_text
    bar_titles = [
        node.text
        for node in svg.findall(
            ".//{http://www.w3.org/2000/svg}g/{http://www.w3.org/2000/svg}title"
        )
    ]
    assert not any(title.startswith(("CRM ARR:", "Billing ARR:")) for title in bar_titles)
    assert not re.search(r"CHF -?[\d,]+\.\d", svg_text)
    for group in owner_groups(data):
        assert group["name"] in text
        assert f'class="owner-id">{group["owner"]}</small>' in text
    assert 'class="owner-heading"' in text and "border-left: 2px solid" in text
    assert any(bar.get("class") == "negative" and float(bar.get("height")) > 0 for bar in bars)
    with proposal.path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        rows = list(reader)
        assert reader.fieldnames[-1] == "approved"
    assert len(rows) == proposal.rows == 8
    assert all(row["approved"] == "" for row in rows)
    workbook = load_workbook(report.workbook_path, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Scorecard",
            "Exceptions",
            "Fixes",
            "Push readiness",
            "ARR bridge",
            "Quarantine",
        ]
        for sheet in workbook:
            assert sheet.freeze_panes == "A2"
            assert all(cell.font.bold for cell in sheet[1])
            assert sheet.auto_filter.ref == sheet.dimensions
            assert all(column.width >= 13 for column in sheet.column_dimensions.values())
            assert not any(cell.data_type == "f" for row in sheet for cell in row)
        exceptions = workbook["Exceptions"]
        headers = [cell.value for cell in exceptions[1]]
        assert len(headers) == len(set(headers))
        assert "finding_reason" in headers
        for column in ("decision", "reason", "assignee"):
            assert all(row[headers.index(column)].value is None for row in list(exceptions)[1:])
        record = next(
            row
            for row in list(exceptions)[1:]
            if row[headers.index("record_id")].value == "PLANT-R04"
        )
        money = record[headers.index("money_at_risk_chf")]
        assert money.value == 18000 and money.data_type == "n" and "CHF" in money.number_format
        assert [cell.value for cell in workbook["Fixes"][1]][-1] == "approved"
        assert all(row[-1].value is None for row in list(workbook["Fixes"])[1:])
        assert isinstance(workbook["Scorecard"]["E2"].value, float)
        assert workbook["Scorecard"]["E2"].number_format == "0.00%"
        assert sum(row[2].value for row in list(workbook["ARR bridge"])[1:]) == -90216
    finally:
        workbook.close()


def test_ranking_unknown_valuation_and_html_escaping(built: dict[str, Path]) -> None:
    data = read_review_data(built["default"])
    source = data.tables["exceptions"].rows[0]
    rows = [
        {**source, "record_id": identifier, "severity": severity, "money_at_risk_chf": amount}
        for identifier, severity, amount in (
            ("small", "leak", Decimal(1)),
            ("high-hygiene", "hygiene", Decimal(10)),
            ("high-leak", "leak", Decimal(10)),
            ("zero", "info", Decimal(0)),
            ("unknown", "leak", None),
        )
    ]
    rows[0]["reason"] = '<script>alert("synthetic")</script>'
    tables = {**data.tables, "exceptions": Table(data.tables["exceptions"].columns, rows)}
    data = replace(data, tables=tables)
    assert [row["record_id"] for row in data.ranked_exceptions] == [
        "high-leak",
        "high-hygiene",
        "small",
        "zero",
        "unknown",
    ]
    assert owner_groups(data)[0]["risk"] == 21
    text = render_html(data)
    assert "&lt;script&gt;" in text and "<script>" not in text
    assert chf(None) == "not valued" and chf(Decimal(0)) == "CHF 0.00"
    summary = quarantine_summary(data)
    assert sum(row["findings"] for row in summary) == 10
    assert all(row["valuation"] == "not valued" for row in summary)


def test_complete_exports_do_not_use_the_show_limit_and_excel_keeps_literal_text(
    built: dict[str, Path], tmp_path: Path
) -> None:
    database = tmp_path / "test.duckdb"
    shutil.copyfile(built["default"], database)
    connection = duckdb.connect(str(database), config={"autoinstall_known_extensions": False})
    try:
        connection.execute(
            "INSERT INTO marts.exceptions SELECT e.check_id,e.severity,e.object_type,"
            "'SYNTHETIC-REVIEW-' || i,e.owner_id,'=2+3',e.details,e.money_at_risk_chf "
            "FROM range(60) n(i), (SELECT * FROM marts.exceptions LIMIT 1) e"
        )
        connection.execute(
            "INSERT INTO marts.fix_proposals SELECT md5('review-' || i),p.check_id,p.object_type,"
            "'SYNTHETIC-FIX-' || i,p.property,p.current_value,p.proposed_value,"
            "p.requires_clear_first,p.reason "
            "FROM range(60) n(i), (SELECT * FROM marts.fix_proposals LIMIT 1) p"
        )
    finally:
        connection.close()
    result = write_report(db_path=database, output_dir=tmp_path)
    proposal = propose_fixes(db_path=database, output_dir=tmp_path)
    assert result.exception_rows == 104 and proposal.rows == 68
    assert "SYNTHETIC-REVIEW-59" in result.html_path.read_text(encoding="utf-8")
    workbook = load_workbook(result.workbook_path, data_only=False)
    try:
        sheet = workbook["Exceptions"]
        headers = [cell.value for cell in sheet[1]]
        literals = [
            row[headers.index("finding_reason")]
            for row in list(sheet)[1:]
            if str(row[headers.index("record_id")].value).startswith("SYNTHETIC-REVIEW-")
        ]
        assert len(literals) == 60
        assert all(cell.value == "=2+3" and cell.data_type == "s" for cell in literals)
    finally:
        workbook.close()


def test_empty_report_and_waterfall_are_valid(built: dict[str, Path], tmp_path: Path) -> None:
    data = read_review_data(built["default"])
    data = replace(
        data,
        crm_arr=Decimal(0),
        billing_arr=Decimal(0),
        tables={name: Table(table.columns, []) for name, table in data.tables.items()},
    )
    chart = waterfall(data)
    assert all(0 <= tick["y"] <= chart["height"] for tick in chart["ticks"])
    assert all(bar["height"] == 0 for bar in chart["bars"])
    text = render_html(data)
    assert "No exceptions." in text and "No quarantined records." in text
    assert "No deals awaiting a billing push." in text


def test_gap_waterfall_filters_zero_categories_sorts_absolute_values_and_chains_correctly(
    built: dict[str, Path],
) -> None:
    data = read_review_data(built["default"])
    chart = waterfall(data)
    start, *bars, end = chart["bars"]
    assert start["label"] == "Start" and start["before"] == start["after"] == 0
    assert end["label"] == "Total gap" and end["amount"] == data.gap
    assert all(bar["amount"] != 0 for bar in bars)
    assert [abs(bar["amount"]) for bar in bars] == sorted(
        (abs(row["amount_chf"]) for row in data.tables["arr_bridge"].rows if row["amount_chf"]),
        reverse=True,
    )
    cumulative = Decimal(0)
    for bar in bars:
        assert bar["before"] == cumulative
        cumulative += bar["amount"]
        assert bar["after"] == cumulative
    assert cumulative == end["after"] == data.gap
    assert all(bar["label"] != "Unexplained" for bar in bars)
    assert "CRM ARR not valued" in {bar["label"] for bar in bars}
    assert all(bar["color"] == "#168064" for bar in bars if bar["amount"] > 0)
    assert all(bar["color"] == "#bc4b2a" for bar in bars if bar["amount"] < 0)


def test_report_uses_all_owner_names_and_sorts_overdue_pushes_first(built: dict[str, Path]) -> None:
    data = read_review_data(built["default"])
    groups = owner_groups(data)
    assert {group["owner"] for group in groups}.issuperset(data.rules.planted_owner_ids)
    assert all(
        group["name"] == data.owner_names[group["owner"]]
        for group in groups
        if group["owner"] in data.owner_names
    )
    push = data.tables["push_readiness"].rows
    assert push[0]["deal_id"] == "PLANT-R04" and push[0]["status"] == "OVERDUE"
    assert push[1]["deal_id"] == "PLANT-R21" and push[1]["status"] == "DRAFT_ONLY"
    assert [data.rules.push_status_order.index(row["status"]) for row in push] == sorted(
        data.rules.push_status_order.index(row["status"]) for row in push
    )


def test_slack_prints_five_findings_and_never_posts_by_default(
    built: dict[str, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("L2C_DB_PATH", str(built["default"]))
    monkeypatch.setenv("SLACK_WEBHOOK_URL", "https://synthetic.invalid/secret")
    post = Mock(side_effect=AssertionError("Unexpected network call"))
    monkeypatch.setattr(requests, "post", post)
    result = CliRunner().invoke(cli.app, ["slack"])
    assert result.exit_code == 0, result.output
    assert "CHF -90,216.00" in result.output and "Scorecard: 27 checks" in result.output
    assert "PLANT-R04: CHF 18,000.00" in result.output
    assert len(re.findall(r"^\d\. ", result.output, re.MULTILINE)) == 5
    assert "change line" not in result.output and "score history" not in result.output
    post.assert_not_called()


def test_slack_send_requires_webhook_and_uses_fake_http(
    built: dict[str, Path], monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("L2C_DB_PATH", str(built["default"]))
    monkeypatch.delenv("SLACK_WEBHOOK_URL", raising=False)
    post = Mock()
    monkeypatch.setattr(requests, "post", post)
    result = CliRunner().invoke(cli.app, ["slack", "--send"])
    assert result.exit_code == 1 and "SLACK_WEBHOOK_URL must be set" in result.output
    post.assert_not_called()
    monkeypatch.setenv("SLACK_WEBHOOK_URL", "https://synthetic.invalid/secret")
    result = CliRunner().invoke(cli.app, ["slack", "--send"])
    assert result.exit_code == 0, result.output
    post.assert_called_once()
    assert post.call_args.kwargs["json"]["text"] == build_digest()
    post.return_value.raise_for_status.assert_called_once()
    post.side_effect = requests.ConnectionError("https://synthetic.invalid/secret")
    with pytest.raises(ValueError, match="Slack delivery failed") as error:
        send_digest("synthetic")
    assert "secret" not in str(error.value)


@pytest.mark.parametrize("command", [["report"], ["fixes", "propose"], ["slack"]])
def test_missing_database_is_a_clear_nonzero_output_error(
    command: list[str], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("L2C_DB_PATH", str(tmp_path / "missing.duckdb"))
    result = CliRunner().invoke(cli.app, command)
    assert result.exit_code == 1 and "run l2c run first" in result.output
    assert not (tmp_path / "missing.duckdb").exists()


def test_empty_slack_and_escaped_mentions(built: dict[str, Path]) -> None:
    data = read_review_data(built["default"])
    tables = {**data.tables, "exceptions": Table(data.tables["exceptions"].columns, [])}
    assert "No exceptions." in format_digest(replace(data, tables=tables))
    row = {**data.tables["exceptions"].rows[0], "reason": "<@synthetic> & <script>"}
    tables["exceptions"] = Table(data.tables["exceptions"].columns, [row])
    message = format_digest(replace(data, tables=tables))
    assert "&lt;@synthetic&gt; &amp; &lt;script&gt;" in message
