"""Self-contained HTML and six-tab Excel snapshots for human review."""

import math
import textwrap
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from time import perf_counter
from uuid import uuid4

from jinja2 import Environment, PackageLoader, StrictUndefined
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from l2c.logging_setup import log_step
from l2c.output_data import ReviewData, Row, Table, atomic_output, chf, chf_whole, read_review_data
from l2c.paths import OUT_DIR, resolve_repo_path


@dataclass(frozen=True)
class ReportOutputs:
    html_path: Path
    workbook_path: Path
    exception_rows: int


def owner_groups(data: ReviewData) -> list[Row]:
    groups: dict[str, list[Row]] = defaultdict(list)
    for row in data.ranked_exceptions:
        groups[row["owner_id"]].append(row)
    result = []
    for owner, rows in groups.items():
        amounts = [row["money_at_risk_chf"] for row in rows if row["money_at_risk_chf"] is not None]
        result.append(
            {
                "owner": owner,
                "name": data.owner_name(owner),
                "rows": rows,
                "risk": sum(amounts) if amounts else None,
            }
        )
    return sorted(
        result, key=lambda group: (group["risk"] is None, -(group["risk"] or 0), group["owner"])
    )


def quarantine_summary(data: ReviewData) -> list[Row]:
    groups: dict[str, list[Row]] = defaultdict(list)
    for row in data.tables["quarantine"].rows:
        groups[row["reason"]].append(row)
    return [
        {
            "reason": reason,
            "findings": len(rows),
            "records": len({(row["source_table"], row["record_id"]) for row in rows}),
            "valuation": "not valued",
        }
        for reason, rows in sorted(groups.items())
    ]


def bridge_label(category: str) -> str:
    if category == "crm_data_gaps":
        return "CRM data gaps (blank amount or dates)"
    return " ".join(
        word.upper() if word.lower() in {"crm", "arr"} else word
        for word in category.replace("_", " ").capitalize().split()
    )


def waterfall(data: ReviewData) -> Row:
    """Keep calculations in Decimal; use floats only for SVG coordinates."""
    bars = [
        {
            "label": "Start",
            "amount": Decimal(0),
            "before": Decimal(0),
            "after": Decimal(0),
            "kind": "total",
        }
    ]
    running = Decimal(0)
    contributions = sorted(
        (row for row in data.tables["arr_bridge"].rows if row["amount_chf"] != 0),
        key=lambda row: (-abs(row["amount_chf"]), row["category"]),
    )
    for row in contributions:
        amount = row["amount_chf"]
        bars.append(
            {
                "label": bridge_label(row["category"]),
                "amount": amount,
                "before": running,
                "after": running + amount,
                "kind": "positive" if amount > 0 else "negative" if amount < 0 else "zero",
            }
        )
        running += amount
    bars.append(
        {
            "label": "Total gap",
            "amount": data.gap,
            "before": Decimal(0),
            "after": data.gap,
            "kind": "total",
        }
    )
    levels = [Decimal(0), *(bar[key] for bar in bars for key in ("before", "after"))]
    low, high = min(levels), max(levels)
    if high == low:
        high = low + Decimal(1)
    span = high - low
    left, top, plot_height, step, bar_width = 120, 52, 290, 112, 74
    colors = {"total": "#173a55", "positive": "#168064", "negative": "#bc4b2a", "zero": "#8998a3"}

    def y(value: Decimal) -> float:
        return top + float((high - value) / span) * plot_height

    for index, bar in enumerate(bars):
        bar.update(
            x=left + index * step,
            color=colors[bar["kind"]],
            y=min(y(bar["before"]), y(bar["after"])),
            height=abs(y(bar["before"]) - y(bar["after"])),
            next_y=y(bar["after"]),
            width=bar_width,
            label_y=top + plot_height + 27,
            amount_y=min(y(bar["before"]), y(bar["after"])) - 10,
            lines=textwrap.wrap(bar["label"], width=16),
        )
    return {
        "bars": bars,
        "width": left + len(bars) * step + 28,
        "height": 455,
        "step": step,
        "baseline": y(Decimal(0)),
        "ticks": [
            {"value": low + span * Decimal(index) / 4, "y": y(low + span * Decimal(index) / 4)}
            for index in range(5)
        ],
    }


def render_html(data: ReviewData) -> str:
    environment = Environment(
        loader=PackageLoader("l2c", "templates"),
        autoescape=True,
        undefined=StrictUndefined,
    )
    environment.filters["chf"] = chf
    environment.filters["chf_whole"] = chf_whole
    environment.filters["percent"] = lambda value: f"{value:.2%}"
    return environment.get_template("report.html.j2").render(
        data=data,
        tables=data.tables,
        owners=owner_groups(data),
        quarantine=quarantine_summary(data),
        chart=waterfall(data),
    )


def _workbook_tables(data: ReviewData) -> dict[str, Table]:
    exceptions = data.tables["exceptions"]
    # Preserve the finding text under its own name so reviewers' reason starts empty.
    exception_columns = tuple(
        "finding_reason" if key == "reason" else key for key in exceptions.columns
    )
    exception_rows = [
        {
            **{("finding_reason" if key == "reason" else key): value for key, value in row.items()},
            "decision": None,
            "reason": None,
            "assignee": None,
        }
        for group in owner_groups(data)
        for row in group["rows"]
    ]
    fixes = data.tables["fix_proposals"]
    quarantine = data.tables["quarantine"]
    return {
        "Scorecard": data.tables["audit_scorecard"],
        "Exceptions": Table((*exception_columns, "decision", "reason", "assignee"), exception_rows),
        "Fixes": Table(
            (*fixes.columns, "approved"), [{**row, "approved": None} for row in fixes.rows]
        ),
        "Push readiness": data.tables["push_readiness"],
        "ARR bridge": Table(
            data.tables["arr_bridge"].columns,
            [
                {**row, "category": bridge_label(row["category"])}
                for row in data.tables["arr_bridge"].rows
            ],
        ),
        "Quarantine": Table(
            (*quarantine.columns, "valuation"),
            [{**row, "valuation": "not valued"} for row in quarantine.rows],
        ),
    }


def _cell_value(value: object) -> object:
    return "; ".join(map(str, value)) if isinstance(value, list | tuple) else value


def _display_value(value: object, column: str) -> str:
    if value is None:
        return ""
    if column.endswith("_chf"):
        return chf(value)
    if column == "pass_rate":
        return f"{value:.2%}"
    return str(value)


def _add_sheet(workbook: Workbook, title: str, table: Table) -> None:
    sheet = workbook.create_sheet(title)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(table.columns)
    for row in table.rows:
        sheet.append([_cell_value(row.get(column)) for column in table.columns])
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173A55")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28
    for index, column in enumerate(table.columns, 1):
        cells = list(sheet.iter_cols(min_col=index, max_col=index))[0]
        width = (
            max(
                len(line)
                for cell in cells
                for line in (
                    (
                        str(cell.value) if cell.row == 1 else _display_value(cell.value, column)
                    ).splitlines()
                    or [""]
                )
            )
            + 3
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(72, max(13, width))
        editable = (title == "Exceptions" and column in ("decision", "reason", "assignee")) or (
            title == "Fixes" and column == "approved"
        )
        for cell in cells[1:]:
            cell.font = Font(name="Arial", size=10, color="172D3E")
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Source text beginning '=' must remain literal text.
            cell.alignment = Alignment(
                horizontal="right" if isinstance(cell.value, int | float | Decimal) else "left",
                vertical="top",
                wrap_text=True,
            )
            if column.endswith("_chf"):
                cell.number_format = '"CHF" #,##0.00;[Red]"CHF" -#,##0.00;"CHF" 0.00'
            elif column == "pass_rate":
                cell.number_format = "0.00%"
            elif isinstance(cell.value, date | datetime):
                cell.number_format = "yyyy-mm-dd"
            elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                cell.number_format = "#,##0"
            if editable:
                cell.fill = PatternFill("solid", fgColor="FFF3CD")
    for row in sheet.iter_rows(min_row=2):
        lines = max(
            sum(
                max(
                    1,
                    math.ceil(
                        len(line) / max(1, sheet.column_dimensions[cell.column_letter].width - 3)
                    ),
                )
                for line in (
                    _display_value(cell.value, table.columns[cell.column - 1]).splitlines() or [""]
                )
            )
            for cell in row
        )
        sheet.row_dimensions[row[0].row].height = min(409, max(22, lines * 14 + 6))


def write_workbook(data: ReviewData, path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        for title, table in _workbook_tables(data).items():
            _add_sheet(workbook, title, table)
        workbook.save(path)
    finally:
        workbook.close()


def write_report(
    *, db_path: str | Path | None = None, output_dir: Path = OUT_DIR, run_id: str | None = None
) -> ReportOutputs:
    started = perf_counter()
    run_id = run_id or uuid4().hex
    result = None
    try:
        data = read_review_data(db_path)
        directory = resolve_repo_path(output_dir)
        html_path, workbook_path = directory / "report.html", directory / "review.xlsx"
        with atomic_output(html_path) as html_temp, atomic_output(workbook_path) as workbook_temp:
            html_temp.write_text(render_html(data), encoding="utf-8")
            write_workbook(data, workbook_temp)
        result = ReportOutputs(html_path, workbook_path, len(data.tables["exceptions"].rows))
        return result
    finally:
        log_step(
            run_id=run_id,
            step="report",
            rows=result.exception_rows if result else None,
            seconds=perf_counter() - started,
            status="ok" if result else "failed",
        )
